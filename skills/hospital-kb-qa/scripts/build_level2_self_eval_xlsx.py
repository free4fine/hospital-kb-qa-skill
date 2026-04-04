#!/usr/bin/env python3
from __future__ import annotations

"""根据标准 JSONL 生成“智慧服务 2 级”自评打分表（XLSX）。"""

import argparse
import json
import re
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True)
BOLD_FONT = Font(bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")


def load_level2_basic_items(jsonl_path: Path) -> list[dict]:
    """从 JSONL 提取 2 级且“基本项=是”的去重条款。"""
    items = []
    seen = set()

    with jsonl_path.open("r", encoding="utf-8") as f:
        for line in f:
            obj = json.loads(line)
            if obj.get("type") != "table_row":
                continue
            if str(obj.get("等级", "")).strip() not in {"2", "2级"}:
                continue
            if str(obj.get("是否为基本项", "")).strip() != "是":
                continue

            project = str(obj.get("业务项目", "")).strip()
            clause = str(obj.get("系统功能评估内容", "")).strip()
            key = (project, clause)
            if not project or not clause or key in seen:
                continue
            seen.add(key)

            m = re.match(r"^（(\d+)）", clause)
            clause_no = m.group(1) if m else ""

            items.append(
                {
                    "序号": str(obj.get("序号", "")).strip(),
                    "类别": str(obj.get("类别", "")).strip(),
                    "业务项目": project,
                    "条款编号": clause_no,
                    "条款内容": clause,
                    "table_index": obj.get("table_index"),
                    "row_index": obj.get("row_index"),
                }
            )

    return items


def apply_table_style(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    """统一设置表格边框、表头样式和单元格对齐。"""
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            cell.border = THIN_BORDER
            if r == min_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_workbook(items: list[dict], output_path: Path) -> None:
    """构建并保存三张工作表：总览、清单、填报说明。"""
    wb = Workbook()

    # Sheet 1：申报总览与关键门槛判断。
    ws = wb.active
    ws.title = "2级申报总览"

    ws["A1"] = "医院智慧服务2级申报自评打分表"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    ws["A3"] = "填报单位"
    ws["B3"] = ""
    ws["A4"] = "评估标准版本"
    ws["B4"] = "医院智慧服务分级评估标准体系（试行）20190801"
    ws["A5"] = "目标等级"
    ws["B5"] = "2级"
    ws["A6"] = "2级最低总分要求"
    ws["B6"] = 20
    ws["A7"] = "2级基本项目要求"
    ws["B7"] = "6项且全部达标"
    ws["A8"] = "2级选择项目要求"
    ws["B8"] = "≥6/11"

    ws["A10"] = "综合评估总分（人工填报）"
    ws["B10"] = ""
    ws["A11"] = "选择项目达标数（人工填报）"
    ws["B11"] = ""
    ws["A12"] = "基本项条款总数"
    ws["B12"] = "=MAX(COUNTA('2级基本项清单'!A:A)-1,0)"
    ws["A13"] = "基本项已达标条款数"
    ws["B13"] = "=COUNTIF('2级基本项清单'!F:F,\"符合\")"
    ws["A14"] = "基本项条款达标率"
    ws["B14"] = "=IF(B12=0,0,B13/B12)"
    ws["A15"] = "基本项目是否全部达标"
    ws["B15"] = "=IF(COUNTIFS('2级基本项清单'!F:F,\"<>\",'2级基本项清单'!F:F,\"<>符合\")=0,\"是\",\"否\")"
    ws["A16"] = "申报建议"
    # 申报建议规则：总分、选择项目、基本项三条件同时满足。
    ws["B16"] = "=IF(AND(B10>=20,B11>=6,B15=\"是\"),\"可申报2级（建议复核）\",\"暂不满足2级申报\")"

    project_counter = OrderedDict()
    for item in items:
        project_counter[item["业务项目"]] = project_counter.get(item["业务项目"], 0) + 1

    ws["A18"] = "2级基本项目（条款数）"
    ws["A18"].font = BOLD_FONT
    row = 19
    for project, cnt in project_counter.items():
        ws.cell(row, 1, project)
        ws.cell(row, 2, cnt)
        row += 1

    for r in range(3, 17):
        ws.cell(r, 1).font = BOLD_FONT
    ws["A18"].font = BOLD_FONT

    for r in range(3, row):
        for c in range(1, 3):
            ws.cell(r, c).border = THIN_BORDER
            ws.cell(r, c).alignment = Alignment(vertical="center", wrap_text=True)

    ws["B14"].number_format = "0.00%"
    ws["B10"].fill = INPUT_FILL
    ws["B11"].fill = INPUT_FILL

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16

    # Sheet 2：逐条检查清单（现场状态 + 证据 + 问题整改）。
    ws2 = wb.create_sheet("2级基本项清单")
    headers = [
        "序号",
        "评估类别",
        "业务项目",
        "条款编号",
        "条款内容",
        "现场状态",
        "自评得分",
        "证据材料/系统截图",
        "问题与差距",
        "整改计划",
        "责任部门",
        "责任人",
        "完成期限",
        "来源定位",
    ]
    ws2.append(headers)

    for i, item in enumerate(items, 1):
        source = f"table={item.get('table_index')}, row={item.get('row_index')}"
        ws2.append(
            [
                i,
                item["类别"],
                item["业务项目"],
                item["条款编号"],
                item["条款内容"],
                "待确认",
                # 状态为“符合”时自动记 1 分，其余状态记 0 分。
                f"=IF(F{i+1}=\"符合\",1,0)",
                "",
                "",
                "",
                "",
                "",
                "",
                source,
            ]
        )

    apply_table_style(ws2, 1, max(1, len(items) + 1), 1, len(headers))
    ws2.freeze_panes = "A2"

    ws2.column_dimensions["A"].width = 7
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 40
    ws2.column_dimensions["D"].width = 10
    ws2.column_dimensions["E"].width = 56
    ws2.column_dimensions["F"].width = 12
    ws2.column_dimensions["G"].width = 10
    ws2.column_dimensions["H"].width = 28
    ws2.column_dimensions["I"].width = 20
    ws2.column_dimensions["J"].width = 24
    ws2.column_dimensions["K"].width = 14
    ws2.column_dimensions["L"].width = 12
    ws2.column_dimensions["M"].width = 12
    ws2.column_dimensions["N"].width = 18

    dv = DataValidation(type="list", formula1='"符合,部分符合,不符合,待确认"', allow_blank=True)
    ws2.add_data_validation(dv)
    dv.add(f"F2:F{max(2, len(items)+1)}")

    # 对“符合/不符合”做颜色提示，便于现场快速浏览。
    ws2.conditional_formatting.add(
        f"F2:F{max(2, len(items)+1)}",
        CellIsRule(operator="equal", formula=['"符合"'], fill=GOOD_FILL),
    )
    ws2.conditional_formatting.add(
        f"F2:F{max(2, len(items)+1)}",
        CellIsRule(operator="equal", formula=['"不符合"'], fill=WARN_FILL),
    )

    # Sheet 3：填报口径与字段说明。
    ws3 = wb.create_sheet("填报说明")
    notes = [
        "使用说明",
        "1. 本表用于医院智慧服务2级申报的内部自评。",
        "2. 先在【2级基本项清单】逐条填写“现场状态”和证据材料。",
        "3. 现场状态口径：符合/部分符合/不符合/待确认。",
        "4. 自评得分列自动计算：符合=1，其余=0。",
        "5. 在【2级申报总览】填写“综合评估总分”“选择项目达标数”。",
        "6. 申报建议为辅助判断，正式申报前应组织专家复核。",
        "7. 2级关键门槛：最低总分20分、基本项目6项全部达标、选择项目至少6/11。",
        "",
        "字段说明",
        "- 来源定位：对应标准抽取记录位置（table_index, row_index），用于审计追溯。",
        "- 证据材料：建议填写制度编号、系统截图路径、报表导出路径等。",
    ]
    for i, text in enumerate(notes, 1):
        ws3.cell(i, 1, text)
    ws3["A1"].font = TITLE_FONT
    ws3.column_dimensions["A"].width = 120

    wb.save(output_path)


def main() -> None:
    """命令行入口。"""
    parser = argparse.ArgumentParser(description="Build Level-2 self-assessment scoring workbook from JSONL.")
    parser.add_argument("--jsonl", type=Path, required=True, help="Input JSONL path")
    parser.add_argument("--output", type=Path, default=Path("医院智慧服务2级申报自评打分表.xlsx"), help="Output XLSX path")
    args = parser.parse_args()

    items = load_level2_basic_items(args.jsonl)
    if not items:
        raise SystemExit("No level-2 basic items found in JSONL.")

    build_workbook(items, args.output)
    print(f"Wrote {args.output} (items={len(items)})")


if __name__ == "__main__":
    main()
